import subprocess
import tempfile
import shutil
from pathlib import Path
from typing import Callable
from translator_utils import is_russian, safe_exists
from handlers_base import BaseHandler

try:
    from docx import Document as DocxDocument
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import openpyxl
    HAS_XLSX = True
except ImportError:
    HAS_XLSX = False

try:
    import pptx
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False

class OfficeHandler(BaseHandler):
    def translate_docx_inplace(self, path: Path) -> bool:
        if not HAS_DOCX: return False
        try:
            doc = DocxDocument(path)
            changed = False
            for p in doc.paragraphs:
                if p.text:
                    new_text = self.translate_text_if_russian(p.text)
                    if new_text and new_text != p.text:
                        p.text = new_text
                        changed = True
            if changed:
                doc.save(path)
                return True
        except Exception as e:
            print(f"  Warning: DOCX error: {path}: {e}")
        return False

    def translate_pptx_inplace(self, path: Path) -> bool:
        if not HAS_PPTX: return False
        try:
            prs = pptx.Presentation(path)
            changed = False
            for slide in prs.slides:
                for shape in slide.shapes:
                    if not shape.has_text_frame: continue
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if run.text:
                                translation = self.translate_text_if_russian(run.text)
                                if translation and translation != run.text:
                                    run.text = translation
                                    changed = True
            if changed:
                prs.save(path)
                return True
        except Exception as e:
            print(f"  Warning: PPTX error: {path}: {e}")
        return False

    def translate_xlsx_inplace(self, path: Path) -> bool:
        if not HAS_XLSX: return False
        try:
            wb = openpyxl.load_workbook(path)
            changed = False
            for sname in wb.sheetnames:
                ws = wb[sname]
                if not hasattr(ws, "iter_rows"): continue
                for row in ws.iter_rows():
                    for cell in row:
                        if cell.__class__.__name__ == "MergedCell": continue
                        v = cell.value
                        if isinstance(v, str):
                            translation = self.translate_text_if_russian(v)
                            if translation is not None and translation != v:
                                ws.cell(row=cell.row, column=cell.column).value = translation
                                changed = True
            if changed:
                wb.save(path)
                wb.close()
                return True
            wb.close()
        except Exception as e:
            print(f"  Warning: XLSX error: {path}: {e}")
        return False

    def upgrade_office_file(self, path: Path, unique_path_func: Callable[[Path], Path]) -> Path:
        mapping = {".doc": ".docx", ".xls": ".xlsx", ".ppt": ".pptx", ".rtf": ".docx", ".odt": ".docx"}
        suf = path.suffix.lower()
        if suf not in mapping or path.name.startswith("~$"): return path
        new_path = path.with_suffix(mapping[suf])
        with self.lock:
            if safe_exists(new_path): new_path = unique_path_func(new_path)
        
        temp_outdir = Path(tempfile.mkdtemp(prefix="soffice_out_"))
        temp_user_dir = tempfile.mkdtemp(prefix="soffice_temp_")
        try:
            proc = subprocess.Popen(["soffice", "--headless", "--norestore", f"-env:UserInstallation=file://{temp_user_dir}", "--convert-to", new_path.suffix[1:], "--outdir", str(temp_outdir), str(path)], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            try:
                proc.wait(timeout=30)
            except subprocess.TimeoutExpired:
                proc.kill()
                return path
            converted_path = Path(temp_outdir) / path.with_suffix(mapping[suf]).name
            if proc.returncode == 0 and safe_exists(converted_path):
                with self.lock:
                    if safe_exists(new_path): new_path = unique_path_func(new_path)
                shutil.move(str(converted_path), str(new_path))
                try: path.unlink()
                except OSError: pass
                return new_path
        except Exception as e: print(f"  Warning: soffice error for {path.name}: {e}")
        finally:
            shutil.rmtree(temp_outdir, ignore_errors=True)
            shutil.rmtree(temp_user_dir, ignore_errors=True)
        return path
