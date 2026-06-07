import subprocess
import tempfile
import shutil
import io
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Callable, Tuple
from .translator_utils import safe_exists
from .handlers_base import BaseHandler

try:
    from docx import Document as DocxDocument

    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import openpyxl

    HAS_XLSX = True
except ImportError:
    HAS_XLSX = False

try:
    import pptx

    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False


class OfficeHandler(BaseHandler):
    def translate_docx_inplace(self, path: Path) -> bool:
        if not HAS_DOCX:
            return False
        try:
            doc = DocxDocument(path)
            changed = False
            for p in doc.paragraphs:
                if p.text:
                    new_text = self.translate_text_if_russian(p.text)
                    if new_text and new_text != p.text:
                        p.text = new_text
                        changed = True
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            if para.text:
                                new_text = self.translate_text_if_russian(para.text)
                                if new_text and new_text != para.text:
                                    para.text = new_text
                                    changed = True
            if changed:
                doc.save(path)
                return True
        except Exception as e:
            print(f"  Warning: DOCX error: {path}: {e}")
        return False

    def translate_pptx_inplace(self, path: Path) -> bool:
        if not HAS_PPTX:
            return False
        try:
            prs = pptx.Presentation(path)
            changed = False
            for slide in prs.slides:
                for shape in slide.shapes:
                    if not shape.has_text_frame:
                        continue
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if run.text:
                                translation = self.translate_text_if_russian(run.text)
                                if translation and translation != run.text:
                                    run.text = translation
                                    changed = True
            for slide in prs.slides:
                if not slide.has_notes_slide:
                    continue
                notes_tf = slide.notes_slide.notes_text_frame
                for para in notes_tf.paragraphs:
                    for run in para.runs:
                        if run.text:
                            translation = self.translate_text_if_russian(run.text)
                            if translation and translation != run.text:
                                run.text = translation
                                changed = True
            if changed:
                prs.save(path)
                return True
        except Exception as e:
            print(f"  Warning: PPTX error: {path}: {e}")
        return False

    def translate_xlsx_inplace(self, path: Path) -> bool:
        if not HAS_XLSX:
            return False
        try:
            wb = openpyxl.load_workbook(path)
            changed = False
            for sname in wb.sheetnames:
                ws = wb[sname]
                if not hasattr(ws, "iter_rows"):
                    continue
                for row in ws.iter_rows():
                    for cell in row:
                        if cell.__class__.__name__ == "MergedCell":
                            continue
                        v = cell.value
                        if isinstance(v, str):
                            translation = self.translate_text_if_russian(v)
                            if translation is not None and translation != v:
                                ws.cell(
                                    row=cell.row, column=cell.column
                                ).value = translation
                                changed = True
            if changed:
                wb.save(path)
                wb.close()
                return True
            wb.close()
        except Exception as e:
            print(f"  Warning: XLSX error: {path}: {e}")
        return False

    def upgrade_office_file(
        self, path: Path, unique_path_func: Callable[[Path], Path]
    ) -> Path:
        mapping = {
            ".doc": ".docx",
            ".xls": ".xlsx",
            ".ppt": ".pptx",
            ".rtf": ".docx",
            ".odt": ".docx",
        }
        suf = path.suffix.lower()
        if suf not in mapping or path.name.startswith("~$"):
            return path
        new_path = path.with_suffix(mapping[suf])
        with self.lock:
            if safe_exists(new_path):
                new_path = unique_path_func(new_path)

        temp_outdir = Path(tempfile.mkdtemp(prefix="soffice_out_"))
        temp_user_dir = tempfile.mkdtemp(prefix="soffice_temp_")
        try:
            proc = subprocess.Popen(
                [
                    "soffice",
                    "--headless",
                    "--norestore",
                    f"-env:UserInstallation=file://{temp_user_dir}",
                    "--convert-to",
                    new_path.suffix[1:],
                    "--outdir",
                    str(temp_outdir),
                    str(path),
                ],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
            )
            try:
                proc.wait(timeout=30)
            except subprocess.TimeoutExpired:
                proc.kill()
                return path
            converted_path = Path(temp_outdir) / path.with_suffix(mapping[suf]).name
            if proc.returncode == 0 and safe_exists(converted_path):
                with self.lock:
                    if safe_exists(new_path):
                        new_path = unique_path_func(new_path)
                shutil.move(str(converted_path), str(new_path))
                try:
                    path.unlink()
                except OSError:
                    pass
                return new_path
        except Exception as e:
            print(f"  Warning: soffice error for {path.name}: {e}")
        finally:
            shutil.rmtree(temp_outdir, ignore_errors=True)
            shutil.rmtree(temp_user_dir, ignore_errors=True)
        return path

    def translate_odf_inplace(self, path: Path) -> bool:
        """Translate OpenDocument files (.odt, .ods, .odp) by editing content.xml in-place."""
        try:
            with zipfile.ZipFile(path, "r") as zf:
                entries = []
                changed = False
                for info in zf.infolist():
                    data = zf.read(info.filename)
                    if info.filename == "content.xml":
                        data, content_changed = self._translate_xml_bytes(data)
                        changed = changed or content_changed
                    entries.append((info, data))
        except Exception as e:
            print(f"  Warning: ODF read error: {path}: {e}")
            return False

        if not changed:
            return False

        try:
            with tempfile.NamedTemporaryFile(
                suffix=path.suffix, dir=path.parent, delete=False
            ) as tmp_fh:
                tmp_path = Path(tmp_fh.name)
            try:
                with zipfile.ZipFile(tmp_path, "w") as out:
                    for info, data in entries:
                        new_info = zipfile.ZipInfo(info.filename)
                        new_info.date_time = info.date_time
                        new_info.compress_type = info.compress_type
                        new_info.comment = info.comment
                        new_info.extra = info.extra
                        new_info.internal_attr = info.internal_attr
                        new_info.external_attr = info.external_attr
                        out.writestr(new_info, data)
                shutil.move(str(tmp_path), str(path))
            except Exception:
                try:
                    tmp_path.unlink(missing_ok=True)
                except OSError:
                    pass
                raise
            return True
        except Exception as e:
            print(f"  Warning: ODF write error: {path}: {e}")
            return False

    def _translate_xml_bytes(self, data: bytes) -> Tuple[bytes, bool]:
        try:
            root = ET.fromstring(data)
        except Exception:
            return data, False

        changed = False
        for node in root.iter():
            if node.text:
                new_text = self.translate_text_if_russian(node.text)
                if new_text is not None and new_text != node.text:
                    node.text = new_text
                    changed = True
            if node.tail:
                new_tail = self.translate_text_if_russian(node.tail)
                if new_tail is not None and new_tail != node.tail:
                    node.tail = new_tail
                    changed = True

        if not changed:
            return data, False

        out = io.BytesIO()
        ET.ElementTree(root).write(out, encoding="utf-8", xml_declaration=True)
        return out.getvalue(), True
